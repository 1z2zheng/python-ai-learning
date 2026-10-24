from pathlib import Path
from zipfile import ZipFile

from docx import Document


path = Path(r"D:\C261008001\output\thesis\基于Streamlit的电竞外设竞品定价与偏好可视化系统本科毕业论文初稿校订版.docx")
doc = Document(path)
text = "\n".join(
    [paragraph.text for paragraph in doc.paragraphs]
    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
)
terms = [
    "它", "了", "不仅", "而且", "他们", "她们", "它们", "她", "他", "此外",
    "第一", "二", "三", "四", "首先", "其次", "最后", "说明", "表示", "反映",
    "：", "“", "”", "（", "）", "(", ")", "—",
]
print({term: text.count(term) for term in terms if text.count(term)})
print({"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "characters": len(text)})
with ZipFile(path) as archive:
    document_xml = archive.read("word/document.xml").decode("utf-8")
    rels_xml = archive.read("word/_rels/document.xml.rels").decode("utf-8")
print({
    "native_equations": document_xml.count("<m:oMath>"),
    "hyperlinks": document_xml.count("<w:hyperlink"),
    "external_relationships": rels_xml.count('TargetMode="External"'),
    "manual_page_breaks": document_xml.count('w:type="page"'),
})
for term in ("了", "二", "三"):
    for line in text.splitlines():
        if term in line:
            print(term, line)
