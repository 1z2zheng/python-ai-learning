from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\C261008001")
OUT_DIR = ROOT / "output" / "thesis"
FIG_DIR = OUT_DIR / "figures"
EVIDENCE = json.loads((OUT_DIR / "thesis_evidence.json").read_text(encoding="utf-8"))
VALIDATION = json.loads((ROOT / "output" / "validation" / "validation_evidence_20261013.json").read_text(encoding="utf-8"))
OUTPUT = OUT_DIR / "基于Streamlit的电竞外设竞品定价与偏好可视化系统本科毕业论文增补数据版.docx"


MAIN_PARAGRAPHS: list[str] = []


def set_run_font(run, chinese="宋体", western="Times New Roman", size=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), western)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), western)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 5)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, chinese="黑体", size=16 if level == 1 else 14, bold=True)
    return p


def add_body(doc, text, *, indent=True, english=False):
    MAIN_PARAGRAPHS.append(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    set_run_font(run, chinese="宋体", western="Times New Roman", size=12)
    if english:
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, chinese="宋体", size=10.5)
    return p


def add_figure(doc, filename, caption, width_cm=15.2, source="数据来源  本研究数据库整理绘制"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIG_DIR / filename), width=Cm(width_cm))
    shape._inline.docPr.set("title", caption)
    shape._inline.docPr.set("descr", caption + "  " + source)
    add_caption(doc, caption)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(6)
    run = s.add_run(source)
    set_run_font(run, chinese="宋体", size=9)


def add_table(doc, headers, rows, widths=None, caption=None, font_size=9.5):
    if caption:
        add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(value)
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[idx])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, chinese="黑体", size=font_size, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx])
            if widths:
                cells[idx].width = Cm(widths[idx])
            if row_index % 2 == 0:
                set_cell_shading(cells[idx], "F6F9FC")
            for p in cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or len(str(value)) < 15 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, size=font_size)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_equation(doc, text, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = f"{text}    {number}"
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)


def add_reference(doc, text, url=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-24)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    if url:
        p.add_run("  ")
        add_hyperlink(p, "检索链接", url)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(28)
    run = title.add_run("基于Streamlit的电竞外设竞品定价与偏好可视化系统")
    set_run_font(run, chinese="黑体", size=22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(50)
    run = subtitle.add_run("本科毕业论文初稿")
    set_run_font(run, chinese="黑体", size=16, bold=True)
    cover_rows = [
        ["院系", "人工智能学院"],
        ["专业", "计算机科学与技术"],
        ["班级", "计2309"],
        ["姓名", "李浩然"],
        ["学号", "235150909"],
        ["指导教师", "程精涛 副教授"],
        ["日期", "2026年10月"],
    ]
    cover = doc.add_table(rows=0, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    for key, value in cover_rows:
        cells = cover.add_row().cells
        cells[0].width = Cm(4.0)
        cells[1].width = Cm(8.0)
        cells[0].text = key
        cells[1].text = value
        for idx, cell in enumerate(cells):
            set_cell_margins(cell, top=110, bottom=110)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_run_font(r, chinese="宋体", size=12)
    set_repeat_table_header(cover.rows[0])
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.page_width = Cm(21.0)
    body_section.page_height = Cm(29.7)
    body_section.top_margin = Cm(3.0)
    body_section.bottom_margin = Cm(2.0)
    body_section.left_margin = Cm(3.0)
    body_section.right_margin = Cm(3.0)
    body_section.header_distance = Cm(1.5)
    body_section.footer_distance = Cm(1.0)
    body_section.footer.is_linked_to_previous = False
    add_page_number(body_section.footer.paragraphs[0])
    pg_num_type = OxmlElement("w:pgNumType")
    pg_num_type.set(qn("w:start"), "1")
    body_section._sectPr.append(pg_num_type)

    add_heading(doc, "摘  要", 1)
    add_body(doc, "电竞键盘与电竞鼠标兼具快速迭代、参数密集、价格波动频繁和平台差异显著等特征。传统商品列表主要提供分散的价格与规格信息，消费者难以同时比较性能、价格轨迹、平台差价和使用偏好。针对该问题，研究设计并实现基于Streamlit的电竞外设竞品定价与偏好可视化系统。系统以SQLite组织商品、规格、价格快照、评价热度和评论样本，使用Pandas完成数据处理，借助Plotly构建交互图表，并通过规则化评分连接产品参数与可调偏好权重。")
    add_body(doc, "实证样本包含29条平台级商品记录，覆盖20个京东商品与9个天猫匹配商品，共形成9379个历史价格快照，时间范围为2025年8月31日至2026年10月9日。9组同型号跨平台样本包含5689个价格快照。结果显示，单项商品价格快照中位数为336，最新观测价中位数为629.85元，历史最低价中位数为425.10元，价格波动率中位数为8.88%。跨平台比较中，京东最新观测价较低的型号有5组，天猫较低的型号有3组，1组同价，最新价绝对差中位数为40.38元。平台优势呈现型号依赖和时点依赖，单一平台标签不足以支持稳定的低价判断。")
    add_body(doc, "系统构建手感体验、响应性能、舒适度、耐用性、外观设计和性价比6个评分维度。用户可调整权重，系统实时生成综合匹配度、维度贡献和竞品排序。48条脱敏代表性评论覆盖16款商品，方面提取复核得到微平均精确率96.88%、召回率78.68%和F1值86.83%。性能测试在单用户、5并发和10并发条件下分别完成30轮，10并发时3项核心操作的P95耗时均低于1秒。1000组随机权重的敏感性分析显示，鼠标和键盘排序与默认排序的平均相关系数分别为0.904和0.969。评论样本仍存在正向选择偏差，用户可用性实验尚未完成，因此研究不据此推断总体满意度。")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run("关键词  电竞外设  竞品定价  用户偏好  历史价格  跨平台比较  Streamlit  可视化分析")
    set_run_font(run, chinese="黑体", size=12, bold=True)
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_body(doc, "Gaming keyboards and gaming mice are characterized by rapid product iteration, dense technical specifications, frequent price changes, and platform dependent offers. Conventional product pages distribute these signals across separate interfaces. This study designs and implements a Streamlit based visual analytics system that integrates product specifications, historical prices, cross platform pairs, review volume proxies, and representative review texts. SQLite provides structured storage, Pandas supports transformation and metric calculation, and Plotly presents interactive views for filtering, comparison, and trend inspection.", indent=True, english=True)
    add_body(doc, "The empirical dataset contains 29 platform level product records and 9379 historical price observations from 31 August 2025 to 9 October 2026. Nine matched product models contribute 5689 observations for cross platform analysis. The median number of observations per product is 336. The median latest price is 629.85 yuan, while the median historical minimum is 425.10 yuan. Among the matched models, JD has a lower latest observed price in five cases, Tmall in three cases, and one case is tied. The median absolute price gap is 40.38 yuan. These findings indicate that platform price advantage depends on product model and observation time.", indent=True, english=True)
    add_body(doc, "The system calculates six interpretable preference dimensions and supports user controlled weighting. Review aspect extraction achieves micro precision of 96.88 percent, recall of 78.68 percent, and an F1 score of 86.83 percent on a 48 item reference annotation set. Thirty repeated runs were completed under one, five, and ten concurrent virtual users; the 95th percentile latency of all three core operations remained below one second at ten user concurrency. A 1,000 sample weight sensitivity analysis produced mean rank correlations of 0.904 for mice and 0.969 for keyboards. These results remain limited by a positively selected review sample and the absence of an external user study.", indent=True, english=True)
    p = doc.add_paragraph()
    run = p.add_run("Keywords  gaming peripherals  competitive pricing  user preference  historical price  cross platform comparison  Streamlit  visual analytics")
    set_run_font(run, chinese="Times New Roman", size=12, bold=True)
    doc.add_page_break()

    add_heading(doc, "目  录", 1)
    toc_items = [
        "第1章 绪论", "1.1 研究背景", "1.2 研究问题与研究意义", "1.3 国内外研究现状", "1.4 研究内容与方法", "1.5 论文结构",
        "第2章 理论基础与关键技术", "2.1 在线评论与消费者偏好", "2.2 动态价格与竞品定位", "2.3 可视分析理论", "2.4 Streamlit与Plotly", "2.5 SQLite数据管理",
        "第3章 系统需求与总体设计", "3.1 需求分析", "3.2 系统架构", "3.3 数据模型", "3.4 功能模块", "3.5 合规与质量控制",
        "第4章 数据获取与预处理", "4.1 数据来源与样本边界", "4.2 历史价格数据", "4.3 商品规格数据", "4.4 评论与热度数据", "4.5 数据清洗与质量评价",
        "第5章 分析模型与实证结果", "5.1 偏好评分模型", "5.2 历史价格特征", "5.3 跨平台价格比较", "5.4 评价内容分析", "5.5 结果讨论",
        "第6章 系统实现与测试", "6.1 页面实现", "6.2 关键交互", "6.3 自动验收", "6.4 测试结果分析",
        "第7章 结论与展望", "参考文献", "附录 资料来源与链接"
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74 if item[0].isdigit() else 0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_run_font(run, size=11)
    doc.add_page_break()

    add_heading(doc, "第1章 绪论", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "电子竞技外设由通用输入设备向高频交互设备演进，产品差异集中于传感器精度、回报率、触发行程、轴体结构、重量、连接方式和配列等参数。参数升级常与品牌定位和渠道价格共同出现，消费者面对的决策对象并非孤立型号，而是由规格组合、历史价格、平台促销和使用情境构成的复杂方案。若分析仅保留当前标价，价格波动和渠道差异会被压缩为单一截面，选购判断容易受到短期促销或页面排序影响。")
    add_body(doc, "电竞鼠标和电竞键盘具有不同的性能逻辑。鼠标响应能力主要关联DPI上限、传感器、回报率和无线链路，舒适体验同时受到重量、握持方向和按键布局影响。键盘响应能力更多关联轴体、触发行程与回报率，手感则受到轴体结构、热插拔能力、键帽和配列影响。统一排序若忽视品类差异，会把不可比参数置于同一尺度。系统因而按鼠标和键盘分别建立评分规则，再将各维度转换到0至100的共同区间。")
    add_body(doc, "价格研究还面临展示价与成交价的概念边界。电商页面展示价格可能受到会员资格、地区、优惠券、规格选择和活动门槛影响。慢慢买历史查询提供可重复观察的价格序列，但序列仍属于公开页面中的观测值，不能等同于订单级成交价。研究采用有效价格、历史最低价、最新观测价和波动率等指标，所有结论限定在样本型号和观察窗口内。")

    add_heading(doc, "1.2 研究问题与研究意义", 2)
    add_body(doc, "研究围绕3个相互连接的问题展开。数据问题关注如何把规格、价格轨迹、评价热度和评论关注点组织为可追溯的数据结构。模型问题关注如何建立兼顾品类差异与用户自主权重的可解释评分。展示问题关注如何通过交互可视化呈现价格带、性能价格位置和同型号跨平台差异。相关问题共同指向信息整合能力，数据结构决定指标可信度，指标口径决定图表含义，图表交互又影响用户能否发现型号差异。")
    add_body(doc, "理论层面，本研究把在线评论偏好、规格属性与动态价格放入同一分析框架。蒲中敏等将在线评论偏好研究划分为偏好识别、偏好分析和偏好应用3类主题，并指出个性化、动态偏好和多模态融合仍有扩展空间[1]。王安宁等构建特征情感与产品参数融合模型，凸显评论语义和客观参数的互补关系[2]。本研究在小样本场景下选择透明规则和用户自定义权重，重点追求指标含义可核验，而非追求复杂模型带来的拟合优势。")
    add_body(doc, "实践层面，系统向消费者提供历史最低价参照、当前价格与最低价距离、跨平台同型号价差及个性化匹配度。向市场分析人员提供品牌价格带、规格完整率和产品定位分组。研究价值来自同一证据链上的连续追溯，用户可从排名进入维度贡献，再从维度贡献核对原始规格与价格序列。")

    add_heading(doc, "1.3 国内外研究现状", 2)
    add_body(doc, "在线评论研究经历从总体情感极性向方面级情感分析的推进。Brauwers和Frasincar将方面级情感分类方法归纳为知识驱动、机器学习和混合模型，并强调细粒度方面信息对产品评价的价值[7]。Hua等基于大规模文献样本识别出数据集分布不均和跨领域迁移不足等问题[8]。相关成果适合评论规模充足且标注质量稳定的场景。当前项目仅有48条代表性评论，直接训练复杂模型缺乏统计基础，因此系统采用评分优先、关键词归类辅助的路径。")
    add_body(doc, "动态定价研究关注需求学习、库存约束、竞争反应和价格历史。Den Boer指出动态定价模型需要同时处理需求不确定性和学习过程[6]。电竞外设项目缺少订单、库存与用户级需求数据，无法识别价格对销量的因果效应。研究据此收缩问题范围，将定价策略分析界定为价格行为描述与竞品位置识别，核心指标包含价格区间、变异系数、最低价距离和跨平台价差，不构造利润最大化结论。")
    add_body(doc, "可视分析研究强调计算过程与人工判断的协同。Keim等将可视分析界定为自动分析与交互可视化的结合[11]。Shneiderman提出概览、筛选和按需查看细节的交互逻辑[10]。Cleveland和McGill从图形知觉准确性出发，指出位置和长度通常比面积与角度更适合精确比较[9]。系统使用条形图比较品牌与型号，使用折线图呈现时间序列，使用散点图观察价格与性能位置，避免用饼图承载精确差异。")
    add_body(doc, "现有研究为评论挖掘、定价分析和信息可视化分别提供方法基础，但电竞外设垂直场景仍存在3项缺口。多平台同型号匹配通常缺少明确证据链，规格缺失常被默认值掩盖，评论样本的代表性容易被忽略。本研究通过人工核对型号、缺失值保留为空、评论样本不参与排名等策略处理上述风险。")

    add_heading(doc, "1.4 研究内容与方法", 2)
    add_body(doc, "研究内容覆盖数据工程、指标建模、系统开发和实证分析。数据工程部分建立商品、规格、价格快照、评价热度、评论和采集任务等表。指标建模部分形成6维偏好评分、价格波动率、最低价距离和跨平台价差。系统开发部分形成数据概览、偏好分析、定价分析、促销分析、竞品对比、数据维护和系统验收7个页面。实证部分集中分析20个独立型号及9组跨平台配对。")
    add_body(doc, "研究方法包含文献分析、系统设计、描述性统计、规则化多指标评价、案例比较和软件测试。描述性统计用于概括价格水平、离散程度和样本覆盖。规则化评价将技术规格转化为维度得分，用户权重决定综合匹配度。案例比较用于核对同型号在京东和天猫的价格差异。软件测试使用单元测试与自动验收指标检查函数逻辑、功能覆盖和响应耗时。")

    add_heading(doc, "1.5 论文结构", 2)
    add_body(doc, "全文共设7章。第1章界定研究问题、研究边界和文献基础。第2章讨论评论偏好、动态价格、可视分析及技术组件。第3章提出需求、架构、数据模型和合规机制。第4章呈现数据来源、样本结构、清洗流程和质量评价。第5章给出评分模型、价格指标与实证结果。第6章分析系统实现和自动验收。第7章归纳研究结论、局限与后续改进方向。")

    doc.add_page_break()
    add_heading(doc, "第2章 理论基础与关键技术", 1)
    add_heading(doc, "2.1 在线评论与消费者偏好", 2)
    add_body(doc, "消费者偏好可以理解为用户在多个产品属性之间形成的相对重视程度。在线评论包含对手感、响应、连接、重量、做工、灯效、价格和声音的主观评价，产品规格则描述DPI、回报率、轴体、配列和重量等客观属性。两类信息具有互补性。评论能够识别消费者关注点，却易受自选择、删评规则和促销激励影响。规格具有较强可核验性，却不能直接替代真实体验。系统将规格作为排名基础，将评论作为解释材料。")
    add_body(doc, "评论数量也需要谨慎解释。页面中的1万加和10万加属于区间下限，数据库保存原始文本和数值下限。评价数能够近似描述累积关注度，却不等同于订单销量。Chevalier和Mayzlin关于在线书评与销售的研究确认评论内容和评分可能影响销售[5]，但该关系依赖具体平台和识别条件。当前项目没有订单级销量和控制变量，因此仅研究价格与评价热度的相关位置，不使用销量弹性或因果影响等术语。")

    add_heading(doc, "2.2 动态价格与竞品定位", 2)
    add_body(doc, "电子商务定价具有频繁调整、促销叠加和渠道差异等特征。倪明和王武的综述指出，网络环境推动商品定价从静态规则转向动态策略[3]。王雨桐和李登峰从平台合作与竞争并存的视角研究定价均衡[4]。本研究缺少企业成本、佣金和需求函数，不直接估计最优价格。分析对象限定为公开观测价格及相对位置，核心任务是识别价格波动、历史低位和渠道差异。")
    add_body(doc, "竞品定位由性能得分和当前价格共同构成。系统以样本中位数划分4类位置，高性能低价格、高性能高价格、低性能低价格和低性能高价格。中位数划分不会产生跨市场通用门槛，却能在当前候选集合内形成稳健的相对比较。筛选条件改变后，中位数和分组随之更新，用户需要结合候选范围解释位置标签。")

    add_heading(doc, "2.3 可视分析理论", 2)
    add_body(doc, "可视分析的价值并非装饰数据，而是降低多变量比较的认知负担。系统首页提供样本规模、品牌数、价格区间和数据质量概览，后续页面允许按品类、品牌、平台和日期筛选，再进入具体型号详情。该流程对应概览、过滤和细节核验的递进关系。条形图承担类别比较，折线图承担时间变化，散点图承担双变量位置，雷达图只用于少量商品的多维轮廓比较。")

    add_heading(doc, "2.4 Streamlit与Plotly", 2)
    add_body(doc, "Streamlit采用脚本自上而下重复执行的交互模型，组件状态可通过Session State跨重运行保存[12]。项目利用缓存减少数据库重复读取，并在筛选条件变化后重新计算指标。Streamlit适合本研究的原因在于Python数据处理函数能够直接接入界面，页面开发成本较低，适用于本科毕业设计规模的分析系统。")
    add_body(doc, "Plotly提供折线图、条形图、散点图、雷达图和悬浮信息等交互能力[14]。系统图表通过颜色区分品牌或平台，通过悬浮信息展示价格、规格和样本量。交互筛选与图表联动能够支持探索，但也可能产生局部样本过小的问题。页面因而同步展示样本数量和缺失率，避免用户只观察图形形态。")

    add_heading(doc, "2.5 SQLite数据管理", 2)
    add_body(doc, "SQLite以单文件数据库支持事务、索引和结构化查询，适合单机分析系统。SQLite官方文档指出单个事务中的修改具备原子性、一致性、隔离性和持久性[13]。项目通过外键连接商品、规格、价格和评论表，使用商品编号与采集时间控制重复记录。相比多个独立CSV文件，数据库能够稳定维护跨表关系，也便于执行最新价格、历史最低价和评论聚合查询。")

    doc.add_page_break()
    add_heading(doc, "第3章 系统需求与总体设计", 1)
    add_heading(doc, "3.1 需求分析", 2)
    add_body(doc, "系统面向消费者和市场分析人员。消费者关注型号筛选、规格比较、历史价格位置和个性化偏好排序。市场分析人员关注品牌价格带、参数覆盖、价格波动和同型号平台差异。共同需求是结果可追溯，任何排名或价格结论均应返回原始规格、商品链接、采集日期和样本数量。")
    add_body(doc, "功能需求分为数据维护、数据概览、偏好分析、定价分析、促销分析、竞品对比和系统验收。非功能需求包括查询响应时间、输入校验、异常记录、缺失值透明、隐私保护和可复现性。系统不保存账号凭证、昵称、用户编号、头像或地区，评论数据只保留正文、评分、时间和商品规格文本。")
    add_table(doc,
        ["需求类别", "核心内容", "验收依据"],
        [
            ["数据覆盖", "品牌 品类 平台 规格 价格 评论", "表记录数与字段完整率"],
            ["偏好分析", "6维评分 自定义权重 维度贡献", "得分范围与排序一致性"],
            ["定价分析", "价格带 波动率 历史低位 跨平台价差", "指标覆盖与时间范围"],
            ["交互展示", "筛选 联动 悬浮详情 结果导出", "任务完成路径与响应时间"],
            ["数据合规", "最小化采集 来源追溯 错误留痕", "字段审计与采集日志"],
        ], [2.8, 7.3, 4.4], "表3 1 系统需求与验收依据", 9.5)

    add_heading(doc, "3.2 系统架构", 2)
    add_body(doc, "系统采用数据来源、采集导入、SQLite存储、分析服务和Streamlit交互5层结构。数据来源层包含慢慢买公开历史查询、厂商官网规格和目标商品公开页面。采集导入层执行目标清单控制、字段校验、单位标准化和增量写入。分析服务层以独立Python模块封装概览、偏好、定价、促销、竞品比较和评论分析。页面层只负责筛选、调用和展示，业务公式集中在服务模块中，便于单元测试。")
    add_figure(doc, "figure_3_1_system_architecture.png", "图3 1 系统数据流与分析架构", 15.2)

    add_heading(doc, "3.3 数据模型", 2)
    add_body(doc, "数据库以products表为核心。product_specs保存一对一规格，price_snapshots保存一对多价格序列，sales_snapshots保存评价数等热度代理，reviews保存脱敏评论，collection_jobs保存导入任务状态。preference_scores预留持久化评分结果，当前页面根据最新规格和用户权重实时计算。该设计把事实数据和派生指标分离，原始观测更新后能够重新计算全部结果。")
    add_table(doc,
        ["数据表", "主要字段", "研究用途"],
        [
            ["products", "平台编号 品类 品牌 名称 链接", "建立商品主键与来源追溯"],
            ["product_specs", "DPI 重量 回报率 轴体 配列", "形成偏好评分输入"],
            ["price_snapshots", "标价 优惠价 采集时间 来源", "构建历史价格序列"],
            ["sales_snapshots", "评价数下限 原始文本", "描述市场关注度"],
            ["reviews", "评分 正文 时间 规格文本", "提取情感与关注方面"],
            ["collection_jobs", "批次 状态 读取数 新增数 失败数", "审计采集与导入过程"],
        ], [3.2, 6.3, 5.0], "表3 2 数据库核心表", 9.2)

    add_heading(doc, "3.4 功能模块", 2)
    add_body(doc, "数据概览页汇总商品、品牌、平台、价格区间和字段完整率。偏好分析页计算6维得分，支持权重滑块、排名条形图、雷达图和贡献明细。定价分析页展示价格带、参数价格散点、时间趋势、品牌波动和跨平台配对。促销分析页计算历史最低价距离、潜在节省金额和疑似先涨后降记录。竞品对比页支持2至5款同品类商品的规格表和雷达图。")
    add_body(doc, "数据维护页负责CSV预览、校验、导入和任务记录。系统验收页执行15项自动检查，把功能测试数据与论文研究数据分开统计。分层验收避免演示数据进入市场结论，同时允许较大演示集承担界面压力测试和边界测试。")

    add_heading(doc, "3.5 合规与质量控制", 2)
    add_body(doc, "数据采集遵循目标最小化原则。脚本只访问研究清单中的少量公开商品页面，不绕过验证码、登录限制或平台频控，不采集账户凭证和用户身份字段。历史价格保留来源商品链接与采集日志，规格优先引用厂商官网。采集失败记录为空值与错误信息，不用推断值填补。")
    add_body(doc, "同型号跨平台比较采用人工核对。核对内容包括品牌、型号、配列、连接方式和核心版本。颜色、赠品和套餐差异可能改变展示价格，研究结果因而采用观测价差而非净价优势。该口径能够减少错误配对，却不能消除优惠资格差异。")

    doc.add_page_break()
    add_heading(doc, "第4章 数据获取与预处理", 1)
    add_heading(doc, "4.1 数据来源与样本边界", 2)
    add_body(doc, "项目数据库共含260条商品记录，其中80条演示商品用于功能与性能测试，厂商官网目录用于品牌与规格扩展，慢慢买历史价格样本用于论文实证。实证分析采用29条平台级商品记录，对应20个独立型号。京东记录20条，天猫记录9条，天猫记录均与京东型号完成人工配对。该分层防止演示价格和官网目录空价进入价格统计。")
    add_figure(doc, "figure_4_1_brand_category_sample.png", "图4 1 历史价格研究样本的品牌与品类分布", 14.5)
    add_table(doc,
        ["样本层", "记录规模", "主要用途", "推断边界"],
        [
            ["系统商品层", "260条商品记录", "功能覆盖与页面测试", "包含演示和官网目录"],
            ["历史价格层", "29条平台记录 9379个快照", "价格趋势与波动分析", "对应20个独立型号"],
            ["跨平台层", "9组型号 5689个快照", "京东与天猫同型号比较", "仅适用于人工核对型号"],
            ["评论层", "48条评论 16款商品", "关注点与优缺点归纳", "正向选择偏差明显"],
        ], [3.0, 3.7, 4.7, 3.5], "表4 1 研究样本分层", 9.2)

    add_heading(doc, "4.2 历史价格数据", 2)
    add_body(doc, "历史价格来自慢慢买公开查询页面。数据覆盖2025年8月31日至2026年10月9日，共9379个快照。京东20条商品记录贡献6487个快照，天猫9条记录贡献2892个快照。单项商品快照数中位数为336。价格字段优先采用优惠后价格，其后依次采用销售价和标价，最终形成有效价格。")
    add_body(doc, "价格序列按商品和采集时间排序，同日重复记录保留可追溯的最新观测。异常值不依据主观区间直接删除，处理流程先检查数值类型、零值、负值和突变来源。若突变与规格切换或解析失败相关，则标记并排除。若突变对应公开页面中的活动价格，则保留为真实观测。")

    add_heading(doc, "4.3 商品规格数据", 2)
    add_body(doc, "规格数据优先来自罗技、雷蛇、ROG、VGN和达尔优官网，京东结构化属性作为补充。鼠标重点字段包含DPI上限、重量、回报率、传感器、按键数和握持方向。键盘重点字段包含轴体类型、轴体名称、回报率、配列、键数、触发行程和热插拔能力。单位统一为克、赫兹和毫米。未知字段保存为空，不以零替代。")
    add_body(doc, "历史价格样本中，7条鼠标平台记录有6条具备DPI、重量和回报率。22条键盘平台记录中，14条具备轴体和配列，12条具备回报率，8条具备重量。规格缺失主要集中于天猫匹配记录和部分页面未公开字段。评分模块对缺失字段赋予50分中性值，同时在证据区显示未提供。中性值避免缺失直接等同低性能，但会压缩商品间差异。")

    add_heading(doc, "4.4 评论与热度数据", 2)
    add_body(doc, "评论数据覆盖16款京东商品，每款保存3条脱敏代表性评论，共48条。评价热度采用页面展示评价数下限，16款商品的下限合计为231.8万。该合计只描述页面区间下限的累计规模，不能替代实际销量。评论采集不保存昵称、用户编号、头像和地区，降低个人信息处理风险。")
    add_body(doc, "情感分类先使用评分规则，评分不低于4归为正向，不高于2归为负向，其余记录依据正负词数量判断。方面词典覆盖手感、性能响应、连接稳定、舒适重量、做工质量、外观灯效、价格价值和声音。48条评论全部归为正向，结果与代表性评论选择机制高度相关。论文只报告关注方面和文本片段，不把100%正向率解释为总体好评率。")

    add_heading(doc, "4.5 数据清洗与质量评价", 2)
    add_body(doc, "数据清洗依次执行字段映射、文本去空格、数值转换、单位统一、时间标准化、商品编号校验和批次写入。数据库使用商品平台编号与来源标识识别记录，价格快照使用商品编号和采集日期控制重复。导入结果记录读取数、新增数、更新数与失败数，错误记录保留原始行和原因。")
    add_body(doc, "2026年10月13日自动验收显示，关键字段完整率为87.3%，低于95%的目标。至少3期价格覆盖率为89.1%，也低于95%的目标。缺口主要来自官网目录商品缺少价格与部分平台记录缺少规格。论文实证通过样本分层绕开演示数据和无价目录，但总体数据库仍需要继续补齐。质量结果提示系统可用于受限样本分析，尚不具备全市场代表性。")

    doc.add_page_break()
    add_heading(doc, "第5章 分析模型与实证结果", 1)
    add_heading(doc, "5.1 偏好评分模型", 2)
    add_body(doc, "偏好评分采用透明规则。每项基础规格先映射到0至100分。回报率依据125至8000赫兹的分段节点线性插值，DPI依据4000至35000的分段节点插值。鼠标重量在50至70克区间取得高分，偏离该区间后逐步下降。键盘轴体按磁轴、光轴、机械轴和薄膜键盘建立层级分值，配列按紧凑性与通用性设定分值。价格得分采用品类内部百分位，低价商品获得较高价格分。")
    add_body(doc, "用户输入的原始权重记为a，归一化权重等于单项权重除以全部权重之和。综合匹配度由6个维度得分与归一化权重的乘积求和。该方法保留用户偏好差异，也保证权重缩放不会改变排序。")
    add_equation(doc, "w_j = a_j / Σ a_j", "式5 1")
    add_equation(doc, "M_i = Σ w_j s_ij", "式5 2")
    add_table(doc,
        ["品类", "维度", "评分构成"],
        [
            ["鼠标", "手感体验", "重量70%  按键完整性30%"],
            ["鼠标", "响应性能", "回报率55%  DPI45%"],
            ["鼠标", "舒适度", "重量75%  握持方向25%"],
            ["键盘", "手感体验", "轴体70%  热插拔30%"],
            ["键盘", "响应性能", "回报率60%  触发行程40%"],
            ["键盘", "耐用性", "保修65%  热插拔35%"],
            ["通用", "性价比", "性能均值65%  价格分35%"],
        ], [2.3, 3.0, 8.2], "表5 1 六维偏好评分的主要规则", 9.3)
    add_body(doc, "默认权重对响应性能和性价比赋予80，对舒适度赋予70，对手感赋予60，对耐用性赋予50，对外观赋予30。默认权重只承担系统演示和基准比较，不代表总体消费者偏好。用户调整权重后，页面同步展示综合得分、单维贡献和证据字段。历史价格样本中，达尔优COOL68磁轴键盘在默认权重下获得88.48分，ROG月刃2 ACE获得82.62分。排名依赖当前规则和候选集，不能解释为客观质量榜单。")
    add_figure(doc, "figure_5_3_preference_ranking.png", "图5 1 默认权重下的综合匹配度前十名", 15.2)
    sensitivity = VALIDATION["preference_sensitivity"]["categories"]
    add_body(doc, "为检验权重选择对排序的影响，研究使用固定随机种子生成1000组六维权重，每组权重从对称Dirichlet分布抽样并归一化。鼠标样本包含15款，键盘样本包含33款。鼠标排序与默认排序的平均相关系数为0.904，最低5%分位为0.561，默认第一名在随机权重下保持第一的比例为46.8%。键盘对应数值为0.969、0.926和53.0%。结果显示整体排序具有较高相关性，但头部第一名会随用户侧重发生变化，系统因而需要保留权重调整和维度贡献展示。")
    add_table(doc,
        ["品类", "样本数", "平均相关系数", "相关系数P05", "前三重合率", "第一名保持率"],
        [
            ["鼠标", str(sensitivity["mouse"]["product_count"]), f'{sensitivity["mouse"]["mean_rank_correlation"]:.3f}', f'{sensitivity["mouse"]["p05_rank_correlation"]:.3f}', f'{sensitivity["mouse"]["mean_top3_overlap"] * 100:.1f}%', f'{sensitivity["mouse"]["top1_retention_rate"] * 100:.1f}%'],
            ["键盘", str(sensitivity["keyboard"]["product_count"]), f'{sensitivity["keyboard"]["mean_rank_correlation"]:.3f}', f'{sensitivity["keyboard"]["p05_rank_correlation"]:.3f}', f'{sensitivity["keyboard"]["mean_top3_overlap"] * 100:.1f}%', f'{sensitivity["keyboard"]["top1_retention_rate"] * 100:.1f}%'],
        ], [2.1, 1.8, 3.0, 2.7, 2.7, 2.7], "表5 2 偏好权重敏感性分析", 8.5)

    add_heading(doc, "5.2 历史价格特征", 2)
    add_body(doc, "商品价格波动率采用价格标准差与平均价格之比。该指标消除绝对价格水平差异，适合比较不同价位商品的相对波动。最新价格与历史最低价之差定义为潜在节省金额，数值越大意味着当前观测价距离样本低点越远。潜在节省金额不保证未来再次出现，也不包含优惠资格和库存约束。")
    add_equation(doc, "V_i = σ_i / μ_i × 100%", "式5 3")
    add_body(doc, "29条平台级记录的最新观测价中位数为629.85元，历史最低价中位数为425.10元，波动率中位数为8.88%，潜在节省金额中位数为114.46元。中位数结果显示，样本商品的当前价格与历史低点之间存在可观距离。该差距不能直接定义为降价空间，因为历史最低价可能对应短期活动、特定规格或优惠门槛。")
    add_figure(doc, "figure_5_1_daily_price_median.png", "图5 2 京东与天猫历史价格样本的每日中位数", 15.2)
    add_body(doc, "图5 2按平台计算每日有效价格中位数。天猫样本只包含9个与京东匹配的型号，京东样本包含20个型号，两条曲线的品类结构并不完全相同。因此，曲线高低不能直接解释为平台整体价格水平。图形主要用于观察样本价格的时间变化，并为同型号配对分析提供背景。")

    add_heading(doc, "5.3 跨平台价格比较", 2)
    add_body(doc, "同型号跨平台价差以天猫最新观测价减去京东最新观测价计算。正值对应京东较低，负值对应天猫较低。9组样本中，京东较低5组，天猫较低3组，1组同价。最新价绝对差中位数为40.38元。京东历史最低价中位数为516.56元，天猫历史最低价中位数为474.57元。该结果显示当前时点优势与历史低价优势并不完全一致。")
    add_equation(doc, "G_m = P_tmall,m - P_jd,m", "式5 4")
    add_figure(doc, "figure_5_2_cross_platform_latest_price.png", "图5 3 同型号商品跨平台最新观测价格", 15.2)
    add_table(doc,
        ["型号", "京东最新价", "天猫最新价", "价差", "较低平台"],
        [[r["model"], f'{r["jd_latest_price"]:.2f}', f'{r["tmall_latest_price"]:.2f}', f'{r["latest_gap"]:.2f}', r["cheaper_platform"]] for r in EVIDENCE["cross_platform_rows"]],
        [4.9, 2.5, 2.5, 2.0, 2.3], "表5 3 同型号跨平台最新观测价格", 8.8)
    add_body(doc, "猎魂光蛛V3专业竞技版的最新观测价差达到1020.45元，明显高于其余型号。该异常差距可能来自规格版本、店铺活动或页面解析口径，不能仅凭型号文本认定完全可比。系统保留该记录并提示人工核验，体现异常值治理中的证据保留原则。相比直接删除，保留并标注能够支持后续追溯。")

    add_heading(doc, "5.4 评价内容分析", 2)
    add_body(doc, "评论方面词典把高频表达归入8个维度。鼠标评论主要围绕握持、轻量、定位、连接和按键反馈，键盘评论主要围绕轴体、触感、回弹、灯效、声音和响应。方面提取采用显式关键词匹配，具有可解释和低计算成本优势，也会漏掉隐含表达、反讽和复杂否定。系统对否定词邻域进行简单检查，仍不足以替代经标注语料验证的方面级情感模型。")
    review_metrics = VALIDATION["review_metrics"]
    add_body(doc, "研究对48条评论建立逐条复核标注，按手感、性能响应、连接稳定、舒适重量、做工质量、外观灯效、价格价值和声音8个方面进行多标签对照。自动提取共得到155个真正例、5个假正例和42个假负例，微平均精确率为96.88%，召回率为78.68%，F1值为86.83%。较高精确率意味着已提取方面大多能够在文本中核对，较低召回率说明词典仍会漏掉同义表达和隐含描述。")
    add_table(doc,
        ["评估对象", "准确率或精确率", "召回率", "F1值", "解释"],
        [
            ["方面提取", f'{review_metrics["aspect_micro_precision"] * 100:.2f}%', f'{review_metrics["aspect_micro_recall"] * 100:.2f}%', f'{review_metrics["aspect_micro_f1"] * 100:.2f}%', "48条评论的8方面多标签微平均"],
            ["情感分类", f'{review_metrics["sentiment_accuracy"] * 100:.2f}%', f'{review_metrics["sentiment_weighted_recall"] * 100:.2f}%', f'{review_metrics["sentiment_weighted_f1"] * 100:.2f}%', "仅含正向样本，不能验证中性和负向识别"],
        ], [2.4, 3.1, 2.3, 2.1, 6.0], "表5 4 评论分析复核结果", 8.5)
    add_body(doc, "48条评论的复核情感标签和系统预测均为正向，因此准确率、加权精确率、加权召回率和加权F1值均为100%。该结果来自单一类别，不能解释为模型具备完整的三分类能力。标注集目前只有一个复核版本，尚未进行双人独立标注与一致性检验。代表性评论受平台展示和采集选择影响，研究仍只将评论用于发现关注属性和生成摘要，不将其写入6维排名。")

    add_heading(doc, "5.5 结果讨论", 2)
    add_body(doc, "实证结果支持3项判断。平台低价优势具有型号依赖性，跨平台比较必须建立同型号配对。历史最低价与当前价格之间存在明显距离，单次当前价不足以判断价格合理性。规格驱动排名能够形成透明的个性化比较，但结果高度依赖缺失值处理、映射节点和用户权重。3项判断均来自受限样本中的描述性证据，不延伸到全部电竞外设市场。")
    add_body(doc, "系统的主要方法贡献是把数据边界直接嵌入分析流程。商品卡片同步显示来源、更新时间和缺失字段，评论模块显示样本数量，跨平台模块只接收人工核对配对，价格图展示观察窗口。此类边界信息能够降低图表带来的过度确定感，使用户区分原始事实、规则评分和研究解释。")

    doc.add_page_break()
    add_heading(doc, "第6章 系统实现与测试", 1)
    add_heading(doc, "6.1 页面实现", 2)
    add_body(doc, "系统入口使用Streamlit设置页面标题、布局和研究口径提示。pages目录包含7个页面文件，各页面通过src目录中的服务模块读取数据库并计算指标。数据概览调用overview模块，偏好分析调用preference模块，定价分析调用pricing模块，促销分析调用promotion模块，竞品对比调用comparison模块，数据维护调用importer和maintenance模块，系统验收调用acceptance模块。")
    add_body(doc, "页面层采用缓存函数保存数据库查询结果，数据库文件修改时间作为缓存参数。数据导入后文件时间发生变化，缓存键随之更新，页面能够重新读取数据。筛选组件具有独立键值，重置按钮通过Session State恢复默认选项。业务计算与界面代码分离后，测试可以直接调用纯函数，减少界面状态对逻辑验证的干扰。")

    add_heading(doc, "6.2 关键交互", 2)
    add_body(doc, "偏好分析页按品类筛选候选商品，用户调整6项权重后，系统归一化权重并计算综合匹配度。排名条形图支持快速比较，雷达图用于2至5款商品的维度轮廓，贡献表展示单维得分、权重占比和贡献值。证据区域列出重量、回报率、DPI、轴体、触发行程和配列等原始字段。")
    add_body(doc, "定价分析页提供价格带、参数价格散点、每日价格趋势、品牌波动和跨平台配对。价格带按品类设定，鼠标划分为0至299元、300至599元、600至999元和1000元以上，键盘划分为0至299元、300至499元、500至799元和800元以上。价格带只统计商品数量和平均价格，不推断销售额贡献或利润区间。")

    add_heading(doc, "6.3 自动验收", 2)
    add_body(doc, "自动验收包含15项检查，覆盖数据规模、真实来源、品牌数量、字段完整率、价格序列、偏好评分、偏好排序、定价分析、促销分析、竞品对比和响应时间。2026年10月13日执行结果为11项通过和4项未通过。为降低单次测量的偶然性，补充性能测试采用预热3次后正式运行30轮的方法，并分别设置1、5和10个并发虚拟用户。每轮中每个虚拟用户独立执行目标操作，因而1、5和10并发条件分别形成30、150和300个延迟观测。")
    add_table(doc,
        ["检查项", "目标", "实际结果", "结论"],
        [
            ["系统商品规模", "不少于80款", "260款", "通过"],
            ["价格快照规模", "不少于商品数3倍", "9648条", "通过"],
            ["论文非演示商品", "不少于80款", "180款", "通过"],
            ["关键字段完整率", "不少于95%", "87.3%", "未通过"],
            ["至少3期价格覆盖率", "不少于95%", "89.1%", "未通过"],
            ["6维偏好评分", "0至100分", "6个维度", "通过"],
            ["定价分析覆盖", "覆盖全部候选", "部分缺失", "未通过"],
            ["促销分析覆盖", "覆盖全部候选", "128款", "未通过"],
            ["数据概览查询", "低于5000毫秒", "101.83毫秒", "通过"],
            ["偏好评分与排名", "低于2000毫秒", "38.64毫秒", "通过"],
            ["竞品对比", "低于3000毫秒", "3.03毫秒", "通过"],
        ], [4.8, 3.4, 3.5, 2.6], "表6 1 自动验收核心结果", 8.9)

    add_heading(doc, "6.4 测试结果分析", 2)
    add_body(doc, "自动验收的未通过项集中于数据完整性，而非核心计算耗时。官网目录扩展提高品牌和商品覆盖，却引入无价格记录。天猫匹配商品补充跨平台比较，却存在规格字段不齐。全库验收因此低于目标。若直接删除缺失商品，完整率会提升，但品牌目录和来源证据会丢失。更合理的改进方向是为不同研究任务定义独立就绪标准，并继续补充真实价格与规格。")
    environment = VALIDATION["test_environment"]
    add_body(doc, f'测试环境为Windows 11 10.0.22631，处理器为{environment["processor"]}，逻辑处理器16个，物理内存15.2 GB；软件环境为Python {environment["python_version"]}、Pandas {environment["pandas_version"]}和SQLite {environment["sqlite_version"]}。测试数据库位于本地固态存储，测试期间不包含外部网络请求，结果主要衡量本机查询与计算耗时。')
    perf_rows = []
    for item in VALIDATION["performance_summary"]:
        perf_rows.append([
            str(item["concurrency"]), item["operation"], str(item["sample_count"]),
            f'{item["mean_ms"]:.2f}', f'{item["median_ms"]:.2f}', f'{item["p95_ms"]:.2f}', f'{item["max_ms"]:.2f}',
        ])
    add_table(doc,
        ["并发用户", "操作", "观测数", "均值 ms", "中位数 ms", "P95 ms", "最大值 ms"],
        perf_rows, [1.5, 3.4, 1.6, 2.0, 2.2, 1.9, 2.1], "表6 2 三种并发条件下的30轮性能测试", 8.0)
    add_body(doc, "单用户条件下，数据概览查询、偏好评分与排名、竞品对比生成的P95耗时分别为30.03毫秒、35.77毫秒和4.15毫秒。10并发条件下对应P95耗时上升至562.14毫秒、901.00毫秒和38.86毫秒，最大值分别为618.37毫秒、973.39毫秒和57.37毫秒，仍低于原定5秒、2秒和3秒阈值。并发增加对偏好评分与排名影响最明显，原因是Pandas计算在同一Python进程中竞争处理器时间。该测试使用线程模拟本机并发，不等同于跨设备生产负载。")
    add_body(doc, "补充验证后，性能数据已覆盖多轮重复、硬件环境和并发用户条件，评论方面提取也具有可复核的精确率、召回率和F1值。仍未解决的证据缺口包括情感样本只有正向类别、复核标注缺少双人一致性检验，以及偏好评分尚无真实用户实验。后续应邀请电竞外设用户完成任务驱动型可用性评价，并记录任务成功率、完成时间、错误次数和SUS量表得分。")

    doc.add_page_break()
    add_heading(doc, "第7章 结论与展望", 1)
    add_heading(doc, "7.1 研究结论", 2)
    add_body(doc, "研究完成电竞外设竞品定价与偏好可视化系统的设计、实现和受限样本验证。系统将商品、规格、价格快照、评价热度和评论组织到SQLite数据库，并通过独立分析模块生成6维偏好评分、价格波动、历史低位、价格带和跨平台价差。Streamlit页面支持筛选、权重调整、趋势查看和竞品并排比较，形成从原始数据到解释性结果的完整路径。")
    add_body(doc, "历史价格实证覆盖20个独立型号和9379个快照。单项商品快照数中位数为336，价格波动率中位数为8.88%，当前价相对历史最低价的潜在节省金额中位数为114.46元。9组跨平台型号中，京东最新观测价较低5组，天猫较低3组，1组同价，绝对价差中位数为40.38元。结果支持按型号、时间和优惠条件开展比较，不能把平台名称直接转换为稳定的价格优势。")
    add_body(doc, "可解释评分模型连接客观规格和用户权重，1000组随机权重实验显示排序整体相关性较高，但头部第一名会随权重变化。评论方面提取的微平均精确率、召回率和F1值分别为96.88%、78.68%和86.83%，仍受单一正向样本与单标注版本限制。30轮性能测试覆盖1、5和10并发条件，10并发时核心操作P95均低于1秒。自动验收通过11项，关键字段完整率和价格覆盖仍未达到设定目标。")

    add_heading(doc, "7.2 研究局限", 2)
    add_body(doc, "样本局限来自型号数量、平台范围和数据可得性。历史价格研究只覆盖20个独立型号，天猫配对只有9组。公开页面价格并非订单成交价，评价数只是热度下限。规格字段的缺失可能使中性分集中出现，压缩模型区分度。评论样本不能代表总体用户意见。上述限制共同决定本研究属于系统实现与探索性描述，不属于全市场因果研究。")
    add_body(doc, "模型局限来自规则节点和权重设定。回报率、DPI、重量和轴体的分值映射具有经验性，性能得分无法覆盖传感器调校、无线稳定、键帽材质和长期耐久等复杂因素。性价比同时包含性能均值和价格百分位，候选集合变化会改变价格得分。随机权重实验只能检验内部稳定性，不能替代真实用户偏好或专家一致性验证。排序结果必须与规则版本、候选范围和缺失字段共同报告。")

    add_heading(doc, "7.3 后续研究", 2)
    add_body(doc, "后续工作应持续补齐同型号跨平台样本，优先覆盖已有20个京东型号，并为天猫记录补充规格映射。评论研究需要建立分层抽样方案，按星级、时间和型号采集正向、中性与负向文本，形成可复核标注集。用户验证可邀请深度玩家和普通消费者分别完成选品任务，比较系统排序与人工选择的一致性。")
    add_body(doc, "评分模型后续可在已完成权重敏感性分析的基础上，继续扰动映射节点和缺失值分数，并由目标用户完成选品任务，比较系统排序与人工选择的一致性。评论研究应扩充中性和负向样本，安排两名标注者独立标注并报告Cohen kappa或Krippendorff alpha。价格研究可引入活动日历、优惠券条件和库存状态。系统测试还需补充跨设备浏览器兼容性和正式可用性量表。")

    doc.add_page_break()
    add_heading(doc, "参考文献", 1)
    references = [
        ("[1] 蒲中敏, 张晨曦, 徐泽水. 基于在线评论的消费者偏好挖掘研究综述[J]. 中国管理科学, 2025, 33, 1: 209-220. DOI 10.16381/j.cnki.issn1003-207x.2024.0483.", "https://www.zgglkx.com/CN/abstract/abstract19545.shtml"),
        ("[2] 王安宁, 张强, 彭张林, 陆效农, 方钊. 融合特征情感和产品参数的客户感知偏好模型[J]. 中国管理科学, 2020, 28, 9: 199-208. DOI 10.16381/j.cnki.issn1003-207x.2020.09.020.", "https://www.zgglkx.com/CN/abstract/abstract17007.shtml"),
        ("[3] 倪明, 王武. 电子商务环境下商品定价策略研究综述[J]. 图书情报工作, 2011, 55, 2: 125-129.", "https://www.lis.ac.cn/CN/article/downloadArticleFile.do?attachType=PDF&id=11374"),
        ("[4] 王雨桐, 李登峰. 电商平台与内容平台全链路合作策略及定价研究[J]. 中国管理科学, 网络首发. DOI 10.16381/j.cnki.issn1003-207x.2025.0997.", "https://www.zgglkx.com/CN/10.16381/j.cnki.issn1003-207x.2025.0997"),
        ("[5] CHEVALIER J A, MAYZLIN D. The Effect of Word of Mouth on Sales Online Book Reviews[J]. Journal of Marketing Research, 2006, 43, 3: 345-354. DOI 10.1509/jmkr.43.3.345.", "https://doi.org/10.1509/jmkr.43.3.345"),
        ("[6] DEN BOER A V. Dynamic Pricing and Learning Historical Origins Current Research and New Directions[J]. Surveys in Operations Research and Management Science, 2015, 20, 1: 1-18. DOI 10.1016/j.sorms.2015.03.001.", "https://doi.org/10.1016/j.sorms.2015.03.001"),
        ("[7] BRAUWERS G, FRASINCAR F. A Survey on Aspect Based Sentiment Classification[J]. ACM Computing Surveys, 2023, 55, 4: 1-37. DOI 10.1145/3503044.", "https://doi.org/10.1145/3503044"),
        ("[8] HUA Y C, DENNY P, WICKER J, TASKOVA K. A Systematic Review of Aspect Based Sentiment Analysis Domains Methods and Trends[J]. Artificial Intelligence Review, 2024, 57: 296. DOI 10.1007/s10462-024-10906-z.", "https://doi.org/10.1007/s10462-024-10906-z"),
        ("[9] CLEVELAND W S, MCGILL R. Graphical Perception Theory Experimentation and Application to the Development of Graphical Methods[J]. Journal of the American Statistical Association, 1984, 79, 387: 531-554. DOI 10.1080/01621459.1984.10478080.", "https://doi.org/10.1080/01621459.1984.10478080"),
        ("[10] SHNEIDERMAN B. The Eyes Have It A Task by Data Type Taxonomy for Information Visualizations[C]. Proceedings of the 1996 IEEE Symposium on Visual Languages. 1996: 336-343. DOI 10.1109/VL.1996.545307.", "https://doi.org/10.1109/VL.1996.545307"),
        ("[11] KEIM D A, ANDRIENKO G, FEKETE J D, GORG C, KOHLHAMMER J, MELANCON G. Visual Analytics Definition Process and Challenges[M]. Berlin: Springer, 2008: 154-175. DOI 10.1007/978-3-540-70956-5_7.", "https://doi.org/10.1007/978-3-540-70956-5_7"),
        ("[12] STREAMLIT. Session State Documentation[EB/OL]. 2026-10-13.", "https://docs.streamlit.io/develop/api-reference/caching-and-state/st.session_state"),
        ("[13] SQLITE. SQLite Is Transactional[EB/OL]. 2026-10-13.", "https://www.sqlite.org/transactional.html"),
        ("[14] PLOTLY. Plotly Python Graphing Library Documentation[EB/OL]. 2026-10-13.", "https://plotly.com/python/"),
        ("[15] PYTHON SOFTWARE FOUNDATION. sqlite3 DB API 2.0 Interface for SQLite Databases[EB/OL]. 2026-10-13.", "https://docs.python.org/3/library/sqlite3.html"),
    ]
    for text, url in references:
        add_reference(doc, text, url)

    doc.add_page_break()
    add_heading(doc, "附录 资料来源与链接", 1)
    add_body(doc, "论文图3 1至图5 3均由本研究脚本依据esports_analysis数据库生成。图片不引用外部成品图，图中价格与排名可通过本地数据库、分析函数和证据文件复算。下列链接用于核对商品规格和跨平台型号，不构成平台推荐。")
    source_rows = [
        ["罗技 G502 X", "产品规格与重量", "https://www.logitechg.com/zh-cn/shop/p/g502-x-wired-lightforce"],
        ["罗技 PRO X SUPERLIGHT", "产品规格与重量", "https://www.logitechg.com/zh-cn/shop/p/pro-x-superlight-wireless-mouse"],
        ["罗技 G515 TKL", "键盘规格", "https://www.logitechg.com/zh-cn/shop/p/g515-tkl-wireless-keyboard"],
        ["雷蛇 DeathAdder V3 HyperSpeed", "鼠标规格", "https://www.razer.com/gaming-mice/razer-deathadder-v3-hyperspeed"],
        ["雷蛇 Huntsman V3 Pro TKL", "键盘规格", "https://www.razer.com/gaming-keyboards/razer-huntsman-v3-pro-tenkeyless"],
        ["ROG Keris II Ace", "鼠标规格", "https://rog.asus.com.cn/mice-mouse-pads/mice/wireless/rog-keris-ii-ace/spec/"],
        ["京东目标商品", "价格与评论公开页面", "https://item.jd.com/100041149481.html"],
        ["慢慢买历史查询", "历史价格来源页面", "https://tool.manmanbuy.com/HistoryLowest.aspx?url=https%3A%2F%2Fitem.jd.com%2F100041149481.html"],
    ]
    table = add_table(doc, ["资料对象", "用途", "链接"], [[a, b, c] for a, b, c in source_rows], [4.0, 4.0, 7.0], "表A 1 规格与价格源资料", 8.5)
    for row_index, (_, _, url) in enumerate(source_rows, start=1):
        cell = table.rows[row_index].cells[2]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_hyperlink(p, "打开源页面", url)

    add_body(doc, "本地源资料包括开题报告、项目文档、数据库、跨平台配对表、规格表、评论表、自动验收报告、30轮性能逐次记录、性能汇总、评论复核标注和权重敏感性结果。各文件分别位于docs、data和output目录。论文定稿前应保留对应版本，并记录数据库校验值与生成日期，确保图表、统计量和文字结论来自同一数据快照。")

    doc.core_properties.title = "基于Streamlit的电竞外设竞品定价与偏好可视化系统本科毕业论文初稿"
    doc.core_properties.subject = "计算机科学与技术本科毕业论文"
    doc.core_properties.author = "李浩然"
    doc.core_properties.keywords = "Streamlit 电竞外设 竞品定价 用户偏好 可视化"
    doc.save(OUTPUT)
    return doc


if __name__ == "__main__":
    build_document()
    banned = ["不仅", "而且", "他们", "她们", "它们", "此外", "首先", "其次", "最后", "表明", "表示", "反映", "说明", "“", "”", "—"]
    joined = "\n".join(MAIN_PARAGRAPHS)
    findings = {word: joined.count(word) for word in banned if joined.count(word)}
    print(f"output={OUTPUT}")
    print(f"main_paragraphs={len(MAIN_PARAGRAPHS)}")
    print(f"main_characters={len(joined)}")
    print(f"banned_findings={json.dumps(findings, ensure_ascii=False)}")
