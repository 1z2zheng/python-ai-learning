from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from docx import Document
from pptx import Presentation


ROOT = Path(r"D:\C261008001")


def docx_text(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return {"paragraphs": paragraphs, "tables": tables}


def pptx_text(path: Path) -> list[dict]:
    prs = Presentation(path)
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        slides.append({"slide": index, "texts": texts})
    return slides


def database_summary(path: Path) -> dict:
    conn = sqlite3.connect(path)
    try:
        tables = [
            row[0]
            for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
            )
        ]
        summary = {}
        for table in tables:
            quoted = table.replace('"', '""')
            summary[table] = conn.execute(f'SELECT COUNT(*) FROM "{quoted}"').fetchone()[0]
        return summary
    finally:
        conn.close()


def main() -> None:
    payload = {
        "opening_reports": {
            path.name: docx_text(path)
            for path in sorted((ROOT / "docs" / "reference").glob("*.docx"))
        },
        "defense_script": docx_text(ROOT / "output" / "C261008001本科毕业设计开题答辩稿.docx"),
        "defense_slides": pptx_text(ROOT / "output" / "C261008001本科毕业设计开题答辩.pptx"),
        "database_counts": database_summary(ROOT / "data" / "esports_analysis.db"),
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
