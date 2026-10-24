import csv
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "thesis"


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.45
    return p


def heading(doc, text, level=2):
    return doc.add_heading(text, level)


def table(doc, headers, rows, widths, caption):
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for i, row in enumerate(t.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EDF1" if i == 0 else ("F5F7F9" if i % 2 else "FFFFFF"))
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), "80")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    el.set(qn(f"w:{key}"), value)
                borders.append(el)
            props.append(borders)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    evidence = json.loads((OUT / "thesis_evidence.json").read_text(encoding="utf-8"))
    review = json.loads((ROOT / "output/review_analysis/review_analysis_summary.json").read_text(encoding="utf-8"))
    sensitivity = read_csv(ROOT / "output/research_revision/score_sensitivity.csv")
    random_results = {}
    for category in ("keyboard", "mouse"):
        rows = [r for r in sensitivity if r["category"] == category and r["scenario"].startswith("random_")]
        random_results[category] = {
            "n": len(rows),
            "mean_correlation": sum(float(r["rank_correlation"]) for r in rows) / len(rows),
            "mean_overlap": sum(float(r["top5_overlap"]) for r in rows) / len(rows),
        }
    feedback = {
        "record_date": "2026-10-14",
        "source": "user_self_report_in_project_conversation",
        "reported_action": "用户报告已阅读全部AI辅助标注，绝大部分判断正确，未填写逐条修改记录",
        "reviewer_identity": "未提供实名",
        "item_level_changes": "not_recorded",
        "all_labels_confirmed": False,
        "independent_double_annotation": False,
        "accuracy": None,
        "status": "reading_completed_disagreements_unrecorded",
    }
    record_path = ROOT / "output/review_analysis/user_reading_review_record.json"
    record_path.write_text(json.dumps(feedback, ensure_ascii=False, indent=2), encoding="utf-8")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    sec.left_margin = sec.right_margin = Inches(0.9)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体" if name == "Normal" else "黑体")
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.add_paragraph("电竞外设竞品分析实证章节修订稿", style="Title")
    body(doc, "本研究以历史报价、商品规格与脱敏评论构建电竞外设分析系统。实证部分集中讨论价格观测边界、评论属性关注结构和规则代理排序的敏感性。研究已完成系统实现及样本描述，独立分类效度、配置一致性与真实用户选品效果仍缺少充分验证。相关结论依据2026年10月14日保存的项目证据，避免将代码执行成功解释为分析模型获得外部验证。")
    heading(doc, "第4章 数据来源与样本处理", 1)
    heading(doc, "4 1 历史报价与评论样本")
    h = evidence["historical_price_sample"]
    body(doc, f"历史报价覆盖{h['start']}至{h['end']}，包含{h['snapshots']:,}条观测，其中京东{h['jd_snapshots']:,}条、天猫{h['tmall_snapshots']:,}条。记录关联40个平台商品对象，配对清单包含20组候选型号。单一型号在两平台形成不同商品对象，型号数量、平台记录数量与时间观测数量应分别报告。日序列中的相邻报价存在时间依赖，价格点数量不能作为独立商品样本规模。")
    table(doc, ["统计对象", "数量", "解释边界"], [
        ["历史报价观测", "13,187", "重复时间观测"], ["平台商品记录", "40", "京东20 天猫20"],
        ["候选型号配对", "20", "配置条件仍待核验"], ["原始评论", "163", "非概率采集样本"],
        ["可分析评论", "140", "通过规则清洗"], ["严格SKU有效评论", "14", "归属及质量均满足条件"],
    ], [2.3, 1.0, 3.3], "表4 1 研究样本统计单位")
    body(doc, "评论来源包含77条历史展示样本和86条评分组定额首页样本，后者集中于三款成功采集的商品。不同型号与品类的评分组分布不一致，混合样本中的情感比例只描述现有文本。评价数作为累积热度代理保存，缺少真实订单销量与统一观察窗口，因此本研究不估计价格对销量的因果效应。")

    heading(doc, "4 2 清洗规则与评论归属")
    body(doc, "清洗保留原始评论，通过规范化字符与空白、型号内文本去重、平台占位评价识别和产品属性筛选生成派生数据。163条原始记录中140条进入规则分析，23条被排除。排除原因允许重叠，同一文本可能同时缺少产品信息并仅涉及物流服务，原因计数不能直接相加解释为排除总数。")
    body(doc, "来源SKU一致的原始评论共19条，其中14条通过清洗。67条来源SKU与目标入口不同，77条历史评论没有来源SKU。来源不同可能涉及颜色、套装或代际变化，现有字段不足以确认全部记录属于同一型号。严格商品级汇总仅采用通过清洗的来源一致记录，其余文本用于发现关注属性和核验线索，禁止直接归因于目标型号的产品质量。")
    heading(doc, "4 3 文本情感与阅读复核")
    body(doc, "方面提取采用手感体验、响应性能、舒适度、耐用性、外观设计和性价比六类词典。情感判断以正文为依据，星级保留为来源属性，不作为方面情感的替代标签。规则采用最长匹配与局部否定检查，避免将没有延迟判断为负向或将不舒服判断为正向。客观重量与连接规格归为中性，同一属性同时具有肯定和批评时保留混合标签。复杂比较、隐含表达与语境依赖仍可能造成误判。")
    body(doc, "另为163条原始评论生成AI辅助标注，覆盖整体情感、六维情感及判断依据。项目使用者报告已阅读全部标签，认为绝大部分判断正确，但没有记录错误条目或最终逐条标签。该过程属于一次阅读复核反馈，尚不足以构成独立标注语料。研究不据此计算准确率、Cohen κ或人工参考标签下的F1，AI标签与规则预测也不混合作为同一统计结果。")
    heading(doc, "4 4 分层统计设计")
    body(doc, "评论按采集方式、日期、评分组、入口型号和SKU归属分别汇总。分层结果有助于识别某些负向体验是否集中于某批商品或某评分组。定额首页采样不具有已知总体抽样概率，缺少各层总体权重时不能还原平台整体满意度。后续模型评价需要冻结词典并保留独立测试文本，避免使用同一批评论同时调整词典并报告识别效果。")

    heading(doc, "第5章 分析模型与实证结果", 1)
    heading(doc, "5 1 评论关注结构")
    body(doc, "140条可分析评论产生365次方面提及。一条评论可评价多个属性，因此提及次数高于评论数量。属性提及率以140条评论为分母，回答某类属性在当前文本中出现的频率，无法直接测量消费者购买决策权重。未提及属性不等于用户不重视该属性，平台展示、写作习惯与使用场景均可能改变文本关注结构。")
    table(doc, ["属性", "提及数", "覆盖率", "正向", "中性", "混合", "负向", "净情感"], [
        [r["dimension"], r["mention_count"], f"{r['mention_rate_pct']:.2f}%", r["positive_count"], r["neutral_count"], r["mixed_count"], r["negative_count"], f"{r['net_sentiment']:.3f}"]
        for r in review["top_dimensions"]
    ], [1.1, .65, .85, .6, .6, .6, .6, .6], "表5 1 当前样本六维文本情感结构")
    body(doc, "净情感值由正向提及数减负向提及数后除以该属性全部提及数，中性与混合均以零贡献进入分母。该指标概括当前规则下的文本极性，混合标签必须同时报告，以免不同体验相互抵消。手感体验覆盖81条评论，占57.86%，关注频率最高。性价比仅有37次提及，较旧版词典减少，主要原因是移除值等过宽匹配，减少与价格无关表达的误归类。")
    body(doc, "舒适度净情感值为0.6863，响应性能为0.6379，耐用性为0.1594，性价比为0.0541。数值只能用于当前规则和抽样结构下的描述，不能转化为品牌质量榜单。键盘与鼠标评论分别受到历史好评展示与新增分层样本的影响，两类净情感差异不能直接解释为产品质量差异。耐用相关文本中的实际故障、售后服务和寿命预期也应逐条区分。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "output/review_analysis/figures/review_dimension_sentiment.png"), width=Inches(5.9))
    doc.add_paragraph("图5 1 当前样本六维文本情感结构", style="Caption").alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading(doc, "5 2 规格代理评分与敏感性")
    body(doc, "规则版本rule-v2.0-proxy将参数评分界定为交互配置、响应规格、重量与配列适配、保修与维护便利、灯效配置和相对价格优势六类代理指标。鼠标手感代理不再重复使用重量，价格优势不再汇入性能均值，键盘维护代理不再重复计入热插拔，灯效代理不再汇入配列。轴体技术名称不承担磁轴优于光轴或机械轴的经验等级。修订降低重复加权，但代理指标与真实体验之间仍需用户实验验证。")
    body(doc, "缺失参数默认占位分为50，属于计算约定，不能解释为真实质量中等。脚本针对研究商品执行1000组随机权重，并分别执行等权、六个单维度消融及缺失占位分25与75的敏感性分析。随机权重结果单独汇总，消融与缺失值情景保留为逐次结果，避免将不同实验设计混作同一总体。")
    table(doc, ["品类", "随机权重次数", "平均排名相关", "前五名平均重合率"], [
        ["键盘" if category == "keyboard" else "鼠标", values["n"], f"{values['mean_correlation']:.3f}", f"{values['mean_overlap']*100:.1f}%"]
        for category, values in random_results.items()
    ], [1.0, 1.4, 1.6, 2.6], "表5 2 随机权重下的排名敏感性")
    body(doc, "排名相关采用商品名次的Spearman相关，前五名重合率比较情景排名与默认排名的头部集合。较高平均相关意味着整体顺序相对接近，不能证明最优商品始终一致，也不能验证默认权重代表真实消费者。鼠标排序的平均相关低于键盘，该候选集对权重调整更敏感。推荐页应提供指标贡献和参数依据，使用户能够理解排序变化。")
    heading(doc, "5 3 历史报价与跨平台比较")
    body(doc, f"40个平台商品记录的末次报价中位数为{h['median_latest_price']:.2f}元，历史最低报价中位数为{h['median_lowest_price']:.2f}元，商品层面波动率中位数为{h['median_volatility_rate']:.2f}%。波动率依据各商品价格标准差与均价计算，商品摘要再采用中位数汇总，减少长序列对整体结果的支配。最低价与末次价之间的差额仅是历史价格位置，不保证用户能够以该最低价完成购买。")
    body(doc, "20组候选配对尚未具备完整配置和优惠核验记录，因此合格平台优势样本为零。系统保留两平台报价与异常提示，配置、店铺、会员条件、补贴资格和报价日期均通过核验后才计入平台优势统计。当前证据不足以支持京东更便宜或天猫更贵的总体结论。")

    heading(doc, "第6章 系统验证与研究限制", 1)
    heading(doc, "6 1 功能验证与数据验收")
    body(doc, "本轮代码验证包含52项测试，覆盖数据导入、规则计算、否定表达、价格维度独立性、缺失值情景隔离和页面执行等行为。首页与七个业务页面均通过执行检查。页面可加载只能证明测试环境下未出现执行异常，图表语义、浏览器交互、下载行为和真实选品任务仍需分别验证。")
    body(doc, "自动验收15项中10项通过，5项未通过。研究规模只统计具有真实京东或天猫历史报价的平台商品记录，无价格目录与演示商品不再用于满足80款目标。三期价格覆盖率以全部目录对象为分母，无历史记录对象计入缺失。口径收紧使缺口更明显，属于验收透明度改进，不代表新增数据损失。")
    table(doc, ["未通过项目", "当前结果", "要求或影响"], [
        ["真实历史价格商品记录", "40条", "目标80款 平台记录非独立型号"],
        ["关键字段完整率", "87.8%", "目标95%"], ["三期价格覆盖率", "46.1%", "目标95% 包含无历史目录"],
        ["定价分析全候选覆盖", "未通过", "无有效价格的目录不能定位"], ["促销指标全目录覆盖", "未通过", "缺少价格历史的对象无指标"],
    ], [2.35, 1.05, 3.2], "表6 1 尚未满足的验收要求")
    heading(doc, "6 2 结论适用范围")
    body(doc, "当前项目能够提供可追溯报价描述、评论关注结构与可解释参数排序。研究贡献主要体现为数据边界进入分析流程，使用户区分商品目录、有效报价、严格SKU评论和待核验候选。现有证据不足以证明消费者总体偏好、品牌耐久差异、实际销量效应或推荐改善真实决策。相关主张必须随样本补充和外部验证逐步建立。")
    body(doc, "后续验证应优先完成分歧标签记录与独立测试语料，逐组核验跨平台配置，再开展任务顺序平衡的用户选品实验。实验计划可邀请12至20名用户，但计划人数不作为已完成样本。记录任务成功、耗时、错误、选中型号与选择理由后，才能评价系统是否帮助用户完成决策。")

    heading(doc, "证据来源与修订衔接", 1)
    body(doc, "本稿涉及的统计量来自项目本地数据和派生结果，不新增外部文献引述。理论基础与文献综述继续沿用原论文中经检索核实的来源，现有实证修订不构成对原参考文献的重新核验。论文合并时需同时更新摘要、数据章节、评价分析、排名图表、系统验收和结论中的旧数字，确保全文使用一致版本。")
    sources = [
        "thesis_evidence.json 保存历史报价样本、字段覆盖和版本信息。",
        "review_analysis_summary.json 与 review_dimension_summary.csv 保存规则清洗与六维统计。",
        "review_stratified_summary.csv 保存日期、入口SKU、评分组和来源归属分层结果。",
        "score_sensitivity.csv 保存随机权重、单维度消融及缺失值情景逐次结果。",
        "cross_platform_verification.json 保存候选配置与报价条件核验状态。",
        "page_execution_checks.json 与 system_acceptance_report.md 保存页面执行和整体验收证据。",
        "评论AI辅助标注待复核.xlsx 保存AI标签与空白人工复核栏。",
        "user_reading_review_record.json 保存用户自述阅读反馈，未声明全部标签正确。",
    ]
    for source in sources:
        body(doc, source)
    heading(doc, "修订时应撤回的旧版结论")
    body(doc, "旧稿中的9组跨平台样本、48条或77条评论、393次方面提及、旧版参数排名与情感准确率均应按当前证据重新处理。163条为原始评论数量，140条为规则分析数量，365次为多标签方面提及。AI辅助标注覆盖原始163条文本，不能以同一数字代替模型独立验证样本。先前平台优势统计缺少配置条件核验，本稿将其调整为待验证候选关系。")
    body(doc, "人工阅读复核已经发生，但缺少逐条分歧清单和最终参考标签。论文可如实记录该过程，不填写未经计数的准确率或一致性指标。后续若取得真实复核记录，应保存修改前后标签与复核日期，再生成对应版本的统计结果。")
    path = OUT / "电竞外设竞品分析实证章节修订稿.docx"
    doc.core_properties.title = "电竞外设竞品分析实证章节修订稿"
    doc.core_properties.author = ""
    for root in (doc.element, doc.styles.element):
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.save(path)
    print(json.dumps({"document":str(path), "random_weight_statistics":random_results, "review_record":str(record_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
